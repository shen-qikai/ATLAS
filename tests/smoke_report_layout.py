"""Approximate font-based visual QA of synthetic reports (not native Excel)."""

from pathlib import Path
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from atlas_pipeline.reports import _display_text, _sizing_font
from test_incremental_pipeline import PipelineIntegrationTests


def main():
    fixture = PipelineIntegrationTests()
    fixture.setUp()
    books = []
    try:
        fixture._write_csv()
        result = fixture._run()
        assert result.processed == 1 and len(result.reports) == 3
        normal = _sizing_font(11, normal=True)
        assert normal is not None, "Visual smoke test needs Windows fonts"
        digit = normal.getlength("0")
        layouts = []
        for filename in result.reports:
            with Path(filename).open("rb") as stream:
                book = load_workbook(stream)
            books.append(book)
            sheet = book.worksheets[0]
            widths = [round(sheet.column_dimensions[get_column_letter(i)].width * digit)
                      for i in range(1, sheet.max_column + 1)]
            heights = [round((sheet.row_dimensions[i].height or 15) * 96 / 72)
                       for i in range(1, sheet.max_row + 1)]
            layouts.append((sheet, widths, heights))
        image = Image.new("RGB", (max(sum(w) for _, w, _ in layouts) + 20,
                                  sum(sum(h) + 38 for _, _, h in layouts) + 10), "white")
        draw = ImageDraw.Draw(image)
        top = 10
        for sheet, widths, heights in layouts:
            draw.text((10, top), sheet.title, font=_sizing_font(11, True), fill="#235789")
            top += 28
            for row, height in zip(sheet.iter_rows(), heights):
                left = 10
                for cell, width in zip(row, widths):
                    fill = "#" + cell.fill.fgColor.rgb[-6:] if cell.fill.fill_type == "solid" else "white"
                    draw.rectangle((left, top, left + width, top + height), fill=fill, outline="#D1D5DB")
                    font = _sizing_font(cell.font.sz or 11, bool(cell.font.bold))
                    text = _display_text(cell)
                    bounds = draw.textbbox((0, 0), text, font=font)
                    assert bounds[2] - bounds[0] <= width - 2, (sheet.title, cell.coordinate, text)
                    assert bounds[3] - bounds[1] <= height, (sheet.title, cell.coordinate)
                    draw.text((left + (width - font.getlength(text)) / 2,
                               top + (height - bounds[3] + bounds[1]) / 2 - bounds[1]),
                              text, font=font, fill="white" if cell.row == 1 else "#111827")
                    left += width
                top += height
            top += 10
        with tempfile.NamedTemporaryFile(prefix="atlas_report_layout_", suffix=".png", delete=False) as target:
            output = Path(target.name)
        image.save(output)
        print(output)
    finally:
        for book in books:
            book.close()
        fixture.tearDown()


if __name__ == "__main__":
    main()
