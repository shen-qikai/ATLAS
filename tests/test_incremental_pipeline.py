"""Synthetic-only integration tests. CSV/XLSX fixtures live in OS temp directories."""

import csv
from dataclasses import replace
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import load_workbook
import yaml

from atlas_pipeline.cleaning import calculate_metrics, merge_and_clean, parse_csv
from atlas_pipeline.config import load_pipeline_profiles
from atlas_pipeline.pipeline import run_pipeline, scan_root, validate_root
from atlas_pipeline.state import product_lock, read_snapshot
from atlas_pipeline.reports import export_reports
from merge_cleaning import WaferMergeToolGUI
from wafer_yield_summary import WaferYieldSummaryApp
from merge_BIN import extract_lot_wafer_name
from preprocessing_core import ProductProfileStore, ProfileConfigError


HEADERS = ["SITE_NUM", "PASSFG", "X_COORD", "Y_COORD", "SOFT_BIN", "VOUT", "UPPER_ONLY", "LOWER_ONLY", "TEXT"]
METADATA = [
    ["", "", "", "", "", "V", "V", "V", ""],
    ["", "", "", "", "", "0", "", "1", ""],
    ["", "", "", "", "", "10", "5", "", ""],
    ["", "", "", "", "", "100", "101", "102", ""],
]
INITIAL_ROWS = [
    [1, "TRUE", 0, 0, 1, 2, 4, 2, "abc"],
    [1, "FALSE", 1, 0, 2, 12, 6, 0, "def"],
]
RETEST_ROWS = [[1, "TRUE", 1, 0, 1, 4, 4, 3, "def"]]


class PipelineIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)
        self.root = self.folder / "FT_DATA"
        self.product = self.root / "LDO" / "WL111"
        self.lot = self.product / "LOT001"
        self.lot.mkdir(parents=True)
        self.config_path = self.folder / "products.yaml"
        self.document = {
            "pipeline_defaults": {"automation": {"test_stage": "CP"}},
            "products": {
                "WL111": {
                    "enabled": True,
                    "naming": {
                        "version": 1,
                        "regex": r"^WL111_(?P<lot>[^_]+)_(?P<wafer>\d+)#_(?P<stage>CP|RT\d+).*\.csv$",
                    },
                    "validation": {"wafer_min": 1, "wafer_max": 25},
                }
            },
        }
        self._save_config()

    def tearDown(self):
        self.temporary.cleanup()

    def test_empty_trailing_columns_ignored_without_changing_raw_or_die_values(self):
        columns = HEADERS + ["", "   "]
        metadata = [row + ["", " "] for row in METADATA]
        rows = [INITIAL_ROWS[0] + ["", ""], INITIAL_ROWS[1]]
        source = self._write_csv(columns=columns, metadata=metadata, rows=rows)
        original = source.read_bytes()
        result = self._run()
        self.assertEqual((result.processed, result.failed), (1, 0))
        self.assertEqual(source.read_bytes(), original)
        record = self._records()[0]
        self.assertEqual(record["metrics"]["record_count"], 2)
        self.assertEqual(record["metrics"]["tests"]["VOUT"]["mean"], 7)
        parsed = parse_csv(source, load_pipeline_profiles(self.config_path)["wl111"])
        self.assertEqual(parsed.columns, HEADERS)

    def test_unnamed_column_with_content_in_metadata_or_last_die_is_rejected(self):
        for metadata_content in (True, False):
            with self.subTest(metadata_content=metadata_content):
                metadata = [row + [""] for row in METADATA]
                rows = [row + [""] for row in INITIAL_ROWS]
                if metadata_content:
                    metadata[0][-1] = "V"
                else:
                    rows[-1][-1] = "0"  # Zero is actual content, not an empty value.
                source = self._write_csv(columns=HEADERS + [""], metadata=metadata, rows=rows)
                with self.assertRaisesRegex(ValueError, "第 10 列.*人工确认"):
                    parse_csv(source, load_pipeline_profiles(self.config_path)["wl111"])

    def test_empty_middle_column_is_not_silently_deleted(self):
        columns = HEADERS[:5] + [""] + HEADERS[5:]
        metadata = [row[:5] + [""] + row[5:] for row in METADATA]
        rows = [row[:5] + [""] + row[5:] for row in INITIAL_ROWS]
        source = self._write_csv(columns=columns, metadata=metadata, rows=rows)
        with self.assertRaisesRegex(ValueError, "非末尾空列名.*第 6 列"):
            parse_csv(source, load_pipeline_profiles(self.config_path)["wl111"])

    def test_duplicate_named_columns_still_rejected_after_trailing_empty_removed(self):
        columns = HEADERS + [" VOUT ", ""]
        source = self._write_csv(columns=columns,
                                 metadata=[row + ["", ""] for row in METADATA],
                                 rows=[row + ["", ""] for row in INITIAL_ROWS])
        with self.assertRaisesRegex(ValueError, "重复列名"):
            parse_csv(source, load_pipeline_profiles(self.config_path)["wl111"])

    def _save_config(self):
        self.config_path.write_text(yaml.safe_dump(self.document), encoding="utf-8")

    def _write_csv(self, wafer=1, suffix="CP", rows=None, metadata=None,
                   columns=None, ending="2026-09-16 10:00:00", lot=None):
        destination = lot or self.lot
        path = destination / f"WL111_{destination.name}_{wafer}#_{suffix}.csv"
        with path.open("w", encoding="utf-8", newline="") as source:
            writer = csv.writer(source)
            if ending:
                writer.writerow(["Ending Time", ending])
            writer.writerow(columns or HEADERS)
            writer.writerows(metadata if metadata is not None else METADATA)
            writer.writerows(rows if rows is not None else INITIAL_ROWS)
        return path

    def _run(self, **kwargs):
        return run_pipeline(self.root, self.config_path, **kwargs)

    def _records(self):
        return list(read_snapshot(self.product)["wafers"].values())

    def _bin_rows(self):
        with (self.product / "WL111_BIN良率.xlsx").open("rb") as source:
            book = load_workbook(source, read_only=True, data_only=True)
            try:
                rows = list(book["BIN汇总"].values)
                return [dict(zip(rows[0], row)) for row in rows[1:]]
            finally:
                book.close()

    def test_first_run_three_reports_cleaned_only_and_raw_immutable(self):
        source = self._write_csv()
        raw = source.read_bytes()
        plan = scan_root(self.root, self.config_path)
        self.assertEqual(plan.products[0].tasks[0].action, "NEW")
        self.assertFalse((self.product / ".atlas").exists())
        result = self._run()
        self.assertEqual((result.processed, result.failed), (1, 0))
        self.assertEqual(len(result.reports), 3)
        self.assertEqual(source.read_bytes(), raw)
        self.assertFalse((self.lot / "summary_data").exists())
        self.assertFalse((self.lot / "summary_clean_onlydata").exists())
        cleaned_path = self.product / self._records()[0]["cleaned_path"]
        with cleaned_path.open(encoding="utf-8", newline="") as cleaned:
            rows = list(csv.reader(cleaned))
        self.assertEqual(len(rows), 7)  # header + four metadata rows + two die rows
        row = self._bin_rows()[0]
        self.assertEqual(row["有效die数量"], 2)
        self.assertEqual(row["总良率"], 0.5)
        self.assertEqual(row["BIN_2"], 0.5)
        self.assertIn(source.name, Path(result.log_path).read_text(encoding="utf-8"))

    def test_report_values_metadata_and_count_formats(self):
        self._write_csv()
        self._run()
        book = load_workbook(self.product / "WL111_测试项失效率.xlsx")
        try:
            sheet = book["测试项失效率"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["VOUT"]).value, "V")
            self.assertEqual(sheet.cell(5, headers["UPPER_ONLY"]).value, 0.5)
            self.assertEqual(sheet.cell(5, headers["LOWER_ONLY"]).value, 0.5)
            self.assertEqual(sheet.cell(5, headers["有效die数量"]).number_format, "0")
            self.assertEqual(sheet.cell(5, headers["VOUT"]).number_format, "0.00%")
        finally:
            book.close()
        book = load_workbook(self.product / "WL111_测试项平均值.xlsx", read_only=True)
        try:
            rows = list(book["测试项均值"].values)
            self.assertEqual(dict(zip(rows[0], rows[4]))["VOUT"], 7)
        finally:
            book.close()

    def test_report_replace_failure_restores_all_previous_files(self):
        self._write_csv()
        self._run()
        reports = list(self.product.glob("*.xlsx"))
        before = {path: path.read_bytes() for path in reports}
        actual_replace = os.replace
        def fail_second(source, target):
            if Path(source).name == "1.xlsx":
                raise PermissionError("Excel lock")
            return actual_replace(source, target)
        snapshot = read_snapshot(self.product)
        with patch("atlas_pipeline.reports.os.replace", side_effect=fail_second):
            with self.assertRaises(RuntimeError):
                export_reports(self.product, snapshot["wafers"])
        self.assertEqual({path: path.read_bytes() for path in reports}, before)

    def test_manual_merge_generates_only_cleaned_output(self):
        self._write_csv()
        gui = object.__new__(WaferMergeToolGUI)
        messages = []
        gui.log = messages.append
        gui.parent = SimpleNamespace(after=lambda delay, callback: callback())
        gui.start_btn = SimpleNamespace(config=lambda **kwargs: None)
        with patch("merge_cleaning.messagebox.showinfo"), patch("merge_cleaning.messagebox.showerror"):
            gui.run_process(str(self.lot), "SITE_NUM", "PASSFG", "X_COORD", "Y_COORD", set(), "1")
        self.assertEqual(len(list((self.lot / "summary_cleaning_data").glob("*.csv"))), 1)
        self.assertFalse((self.lot / "summary_data").exists())
        self.assertFalse((self.lot / "summary_clean_onlydata").exists())

    def test_cleaned_output_is_compatible_with_existing_yield_and_bin_readers(self):
        self._write_csv()
        self._run()
        cleaned = self.product / self._records()[0]["cleaned_path"]
        app = object.__new__(WaferYieldSummaryApp)
        row, meta, columns = app.process_single_param_file((cleaned.name, str(cleaned.parent), "YieldRate", "TotalCount"))
        self.assertEqual(row["VOUT"], 0.5)
        self.assertEqual(meta["VOUT"], ["V", 0, 10])
        row, bins = app.process_single_bin_file((cleaned.name, str(cleaned.parent), "SOFT_BIN"))
        self.assertEqual(row["BIN_1"], 1)
        self.assertEqual(row["_total_"], 2)
        row, headers, columns = app.process_single_avg_file((cleaned.name, str(cleaned.parent)))
        self.assertEqual(row["VOUT"], 7)
        self.assertEqual(row["Lot_WaferID"], "LOT001_W01")
        self.assertEqual(extract_lot_wafer_name(cleaned.name), "LOT001_W1")

    def test_unchanged_run_does_not_recalculate_or_rewrite_reports(self):
        self._write_csv()
        self._run()
        report = self.product / "WL111_BIN良率.xlsx"
        modified = report.stat().st_mtime_ns
        with patch("atlas_pipeline.pipeline.merge_and_clean", side_effect=AssertionError("must skip")), \
             patch("atlas_pipeline.pipeline.file_hash", side_effect=AssertionError("must reuse fingerprint")):
            result = self._run()
        self.assertEqual((result.processed, result.skipped, result.failed), (0, 1, 0))
        self.assertEqual(result.reports, [])
        self.assertEqual(report.stat().st_mtime_ns, modified)

    def test_new_wafer_only_added_once(self):
        self._write_csv()
        self._run()
        self._write_csv(wafer=2)
        result = self._run()
        self.assertEqual((result.processed, result.skipped), (1, 1))
        self.assertEqual(len(self._bin_rows()), 2)
        self._run()
        self.assertEqual(len(self._bin_rows()), 2)

    def test_retest_replaces_wafer_metrics_without_duplicate_row(self):
        self._write_csv()
        self._run()
        self._write_csv(suffix="RT1", rows=RETEST_ROWS, ending="2026-09-16 11:00:00")
        self.assertEqual(scan_root(self.root, self.config_path).products[0].tasks[0].action, "INPUT_CHANGED")
        result = self._run()
        self.assertEqual(result.processed, 1)
        rows = self._bin_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["总良率"], 1)
        metrics = self._records()[0]["metrics"]
        self.assertEqual(metrics["record_count"], 2)
        self.assertEqual(metrics["tests"]["VOUT"]["mean"], 3)

    def test_corrected_source_recalculates(self):
        self._write_csv()
        self._run()
        self._write_csv(rows=RETEST_ROWS)
        self.assertEqual(self._run().processed, 1)
        self.assertEqual(self._bin_rows()[0]["有效die数量"], 1)

    def test_mtime_only_change_with_same_content_skips(self):
        source = self._write_csv()
        self._run()
        stat = source.stat()
        os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1000000000))
        self.assertEqual(self._run().processed, 0)

    def test_complete_hash_verification_detects_same_size_same_mtime_change(self):
        source = self._write_csv()
        self._run()
        stat = source.stat()
        before = source.read_bytes()
        after = before.replace(b",2,4,2,abc", b",3,4,2,abc")
        self.assertNotEqual(before, after)
        source.write_bytes(after)
        os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        self.assertEqual(self._run(verify_hashes=True).processed, 1)

    def test_report_deletion_recovers_from_cache_without_raw_calculation(self):
        self._write_csv()
        self._run()
        (self.product / "WL111_BIN良率.xlsx").unlink()
        with patch("atlas_pipeline.pipeline.merge_and_clean", side_effect=AssertionError("must use cache")):
            result = self._run()
        self.assertEqual((result.processed, result.skipped), (0, 1))
        self.assertEqual(len(result.reports), 3)

    def test_cleaned_deletion_reprocesses_affected_wafer(self):
        self._write_csv()
        self._run()
        (self.product / self._records()[0]["cleaned_path"]).unlink()
        self.assertEqual(self._run().processed, 1)

    def test_failed_retest_keeps_old_cleaned_but_excludes_stale_metrics_then_retries(self):
        self._write_csv()
        self._run()
        cleaned_path = self.product / self._records()[0]["cleaned_path"]
        old_cleaned = cleaned_path.read_bytes()
        bad = self._write_csv(suffix="RT1", rows=RETEST_ROWS)
        bad.write_text("bad csv", encoding="utf-8")
        result = self._run()
        self.assertGreater(result.failed, 0)
        self.assertEqual(self._records()[0]["status"], "failed")
        self.assertEqual(cleaned_path.read_bytes(), old_cleaned)
        self.assertEqual(self._bin_rows(), [])
        self._write_csv(suffix="RT1", rows=RETEST_ROWS, ending="2026-09-16 11:00:00")
        self.assertEqual(self._run().processed, 1)
        self.assertEqual(self._bin_rows()[0]["总良率"], 1)

    def test_missing_input_is_explicitly_unavailable(self):
        source = self._write_csv()
        self._write_csv(wafer=2)
        self._run()
        source.unlink()
        result = self._run()
        self.assertEqual(result.unavailable, 1)
        self.assertEqual(len(self._bin_rows()), 1)
        self.assertIn("missing", [record["status"] for record in self._records()])

    def test_partial_input_deletion_does_not_recompute_from_remaining_retest(self):
        initial = self._write_csv()
        self._write_csv(suffix="RT1", rows=RETEST_ROWS, ending="2026-09-16 11:00:00")
        self._run()
        cleaned_path = self.product / self._records()[0]["cleaned_path"]
        before = cleaned_path.read_bytes()
        initial.unlink()
        result = self._run(force=True)
        self.assertEqual(result.processed, 0)
        self.assertGreater(result.failed, 0)
        self.assertEqual(cleaned_path.read_bytes(), before)
        self.assertEqual(self._bin_rows(), [])

    def test_numeric_metadata_pass_fail_limits_are_not_data_rows(self):
        metadata = [list(row) for row in METADATA]
        metadata[1][1] = "0"
        metadata[2][1] = "1"
        self._write_csv(metadata=metadata)
        self.assertEqual(self._run().processed, 1)

    def test_config_change_forces_recalculation_and_count_warning_not_rejection(self):
        self._write_csv()
        self._run()
        self.document["products"]["WL111"]["wafer"] = {"expected_die_count": 100, "die_count_tolerance": 50}
        self._save_config()
        result = self._run()
        self.assertEqual((result.processed, result.failed), (1, 0))
        self.assertEqual(self._bin_rows()[0]["数量状态"], "数量偏差警告")

    def test_die_count_within_tolerance_is_not_failure(self):
        self.document["products"]["WL111"]["wafer"] = {"expected_die_count": 52, "die_count_tolerance": 50}
        self._save_config()
        self._write_csv()
        self.assertEqual(self._run().failed, 0)
        self.assertEqual(self._bin_rows()[0]["数量状态"], "容差内")

    def test_new_columns_and_limit_versions_are_not_lost(self):
        self._write_csv()
        self._run()
        metadata = [row + [value] for row, value in zip(METADATA, ("V", "0", "10", "103"))]
        metadata[2][5] = "11"
        self._write_csv(wafer=2, columns=HEADERS + ["NEW_TEST"], metadata=metadata,
                        rows=[row + [3] for row in INITIAL_ROWS])
        self.assertEqual(self._run().processed, 1)
        book = load_workbook(self.product / "WL111_测试项失效率.xlsx", read_only=True)
        try:
            rows = list(book["测试项失效率"].values)
            headers = rows[0]
            self.assertIn("NEW_TEST", headers)
            self.assertEqual(rows[3][headers.index("VOUT")], "见测试项规格")
            self.assertEqual(rows[4][headers.index("NEW_TEST")], "NA")
            self.assertEqual(rows[5][headers.index("NEW_TEST")], 0)
        finally:
            book.close()

    def test_force_and_cache_only_modes(self):
        self._write_csv()
        self._run()
        self.assertEqual(self._run(force=True).processed, 1)
        with patch("atlas_pipeline.pipeline._fingerprint", side_effect=AssertionError("cache only")):
            result = self._run(rebuild_reports=True)
        self.assertEqual(result.processed, 0)
        self.assertEqual(len(result.reports), 3)

    def test_excel_failure_retry_does_not_recalculate_successful_wafer(self):
        self._write_csv()
        with patch("atlas_pipeline.reports.export_reports", side_effect=RuntimeError("Excel is open")):
            result = self._run()
        self.assertEqual((result.processed, result.failed), (1, 1))
        snapshot = read_snapshot(self.product)
        self.assertNotEqual(snapshot["generation"], snapshot["report_generation"])
        with patch("atlas_pipeline.pipeline.merge_and_clean", side_effect=AssertionError("already cached")):
            result = self._run()
        self.assertEqual((result.processed, result.skipped, result.failed), (0, 1, 0))
        self.assertEqual(len(result.reports), 3)

    def test_unmatched_file_blocks_whole_lot_not_partial_success(self):
        self._write_csv()
        (self.lot / "unknown.csv").write_text("unknown", encoding="utf-8")
        result = self._run()
        self.assertEqual(result.processed, 0)
        self.assertGreater(result.failed, 0)
        self.assertEqual(self._bin_rows(), [])
        self.assertFalse((self.lot / "summary_cleaning_data").exists())

    def test_lot_folder_mismatch_is_not_guessed(self):
        source = self._write_csv()
        source.rename(self.lot / "WL111_OTHERLOT_1#_CP.csv")
        plan = scan_root(self.root, self.config_path)
        self.assertTrue(plan.products[0].errors)

    def test_dotted_lot_is_supported(self):
        lot = self.product / "LOT001.1"
        lot.mkdir()
        self._write_csv(lot=lot)
        self.assertEqual(self._run().processed, 1)

    def test_vendor_filename_prefix_can_differ_from_product_folder(self):
        self.document["products"]["WL111"]["naming"]["regex"] = (
            r"^VENDORX-CP_(?P<lot>[^_]+)_(?P<wafer>\d+)#_(?P<stage>CP|RT\d+)\.csv$"
        )
        self._save_config()
        source = self._write_csv()
        source.rename(self.lot / "VENDORX-CP_LOT001_1#_CP.csv")
        self.assertEqual(self._run().processed, 1)
        self.assertEqual(self._bin_rows()[0]["Product"], "WL111")

    def test_unknown_product_does_not_block_known_product(self):
        self._write_csv()
        (self.root / "DCDC" / "UNKNOWN" / "LOT_X").mkdir(parents=True)
        result = self._run()
        self.assertEqual(result.processed, 1)
        self.assertGreater(result.failed, 0)

    def test_nonoverlapping_split_files_without_times_are_valid(self):
        self._write_csv(suffix="CP_part1", rows=INITIAL_ROWS[:1], ending=None)
        self._write_csv(suffix="CP_part2", rows=INITIAL_ROWS[1:], ending=None)
        self.assertEqual(self._run().processed, 1)
        self.assertEqual(self._bin_rows()[0]["有效die数量"], 2)

    def test_overlapping_split_files_with_ambiguous_order_fail(self):
        self._write_csv(suffix="CP_part1", ending=None)
        self._write_csv(suffix="CP_part2", ending=None)
        self.assertGreater(self._run().failed, 0)
        self.assertEqual(self._bin_rows(), [])

    def test_single_sided_limits_and_average_values(self):
        source = self._write_csv()
        profile = load_pipeline_profiles(self.config_path)["wl111"]
        metrics = calculate_metrics(merge_and_clean([source], profile), profile)
        self.assertEqual(metrics["tests"]["UPPER_ONLY"]["passed_count"], 1)
        self.assertEqual(metrics["tests"]["LOWER_ONLY"]["passed_count"], 1)
        self.assertEqual(metrics["tests"]["VOUT"]["mean"], 7)
        self.assertEqual(metrics["tests"]["TEXT"]["mean"], None)

    def test_rt_rule_matches_existing_initial_pass_plus_all_retest_semantics(self):
        initial = self._write_csv()
        rt = self._write_csv(suffix="RT1", rows=RETEST_ROWS)
        profile = replace(load_pipeline_profiles(self.config_path)["wl111"], cleaning_mode="rt")
        cleaned = merge_and_clean([initial, rt], profile)
        self.assertEqual(len(cleaned.data), 2)
        self.assertTrue((cleaned.data["PASSFG"] == "TRUE").all())

    def test_blank_unit_row_is_preserved(self):
        metadata = [[""] * len(HEADERS)] + METADATA[1:]
        source = self._write_csv(metadata=metadata)
        profile = load_pipeline_profiles(self.config_path)["wl111"]
        self.assertEqual(len(merge_and_clean([source], profile).metadata), 4)

    def test_source_repo_is_rejected(self):
        with self.assertRaises(ValueError):
            validate_root(Path(__file__).resolve().parents[1])

    def test_parallel_product_lock_rejects_second_writer(self):
        with product_lock(self.product):
            with self.assertRaises(RuntimeError):
                with product_lock(self.product):
                    pass

    def test_duplicate_yaml_products_is_rejected(self):
        self.config_path.write_text("products: {}\nproducts: {}\n", encoding="utf-8")
        with self.assertRaises(ProfileConfigError):
            ProductProfileStore.from_yaml(self.config_path)


if __name__ == "__main__":
    unittest.main()
