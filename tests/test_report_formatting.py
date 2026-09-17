"""Synthetic formatting checks: fit displayed values without changing metrics."""

from copy import deepcopy
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import test_incremental_pipeline as fixtures
from atlas_pipeline.reports import (
    _bin_book, _display_text, _fit_columns, _sizing_font, _style, _test_book, _text_pixels,
)
from atlas_pipeline.state import read_snapshot


class ReportFormattingTests(unittest.TestCase):
    def test_all_cells_centered_and_short_columns_compact(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Wafer", "BIN_1", "有效die数量", "测试项"])
        sheet.append(["W01", 0.5, 10000, 1 / 3])
        _style(sheet, percent_columns=(2,), numeric_columns=(4,))
        for row in sheet:
            for cell in row:
                self.assertEqual((cell.alignment.horizontal, cell.alignment.vertical), ("center", "center"))
        self.assertLess(sheet.column_dimensions["A"].width, 14)
        self.assertLess(sheet.column_dimensions["B"].width, 14)
        self.assertEqual(sheet.row_dimensions[1].height, 20)
        self.assertEqual(sheet.sheet_format.defaultRowHeight, 15)
        self.assertEqual(sheet["B2"].value, 0.5)
        self.assertEqual(sheet["B2"].number_format, "0.00%")
        self.assertEqual(sheet["D2"].value, 1 / 3)
        self.assertEqual(sheet["D2"].number_format, "0.000000")

    def test_measures_displayed_precision_not_raw_floats(self):
        sheet = Workbook().active
        sheet.append([1 / 3, 1 / 3, 123, "NA"])
        sheet["A1"].number_format = "0.00%"
        sheet["B1"].number_format = "0.000000"
        sheet["C1"].number_format = "0"
        self.assertEqual([_display_text(c) for c in sheet[1]], ["33.33%", "0.333333", "123", "NA"])

    def test_all_rows_including_after_40_and_chinese_headers_are_measured(self):
        sheet = Workbook().active
        sheet.append(["中文测试项名称", "Name"])
        for _ in range(50):
            sheet.append(["值", "short"])
        label = "A_very_long_late_row_" * 4
        sheet.append(["值", label])
        _style(sheet)
        font = _sizing_font(11, normal=True)
        digit = font.getlength("0") if font is not None else 7
        self.assertGreaterEqual(sheet.column_dimensions["A"].width * digit,
                                _text_pixels("中文测试项名称", 11, True) + 6)
        self.assertGreaterEqual(sheet.column_dimensions["B"].width * digit,
                                _text_pixels(label, 10) + 6)
        self.assertGreater(sheet.column_dimensions["B"].width, 32)
        self.assertFalse(sheet["B52"].alignment.wrap_text)

    def test_long_and_multiline_content_wraps_only_when_necessary(self):
        sheet = Workbook().active
        sheet.append(["说明", "短列"])
        sheet.append(["路径" * 500, "W01"])
        sheet.append(["first\nsecond", "W02"])
        _style(sheet)
        self.assertEqual(sheet.column_dimensions["A"].width, 255)
        self.assertTrue(sheet["A2"].alignment.wrap_text)
        self.assertTrue(sheet["A3"].alignment.wrap_text)
        self.assertFalse(sheet["B2"].alignment.wrap_text)
        self.assertGreater(sheet.row_dimensions[2].height, 15)
        self.assertGreaterEqual(sheet.row_dimensions[3].height, 30)
        self.assertEqual(sheet["A2"].value, "路径" * 500)

    def test_missing_fonts_fallback_sizes_chinese_and_stays_centered(self):
        sheet = Workbook().active
        sheet.append(["晶圆总die数量", "Wafer"])
        sheet.append([123, "W01"])
        _text_pixels.cache_clear()
        with patch("atlas_pipeline.reports._sizing_font", return_value=None):
            _style(sheet)
        self.assertGreater(sheet.column_dimensions["A"].width, sheet.column_dimensions["B"].width)
        self.assertEqual(sheet["A2"].alignment.horizontal, "center")
        _text_pixels.cache_clear()

    def test_three_saved_reports_centered_and_cache_only_rebuild_preserves_metrics(self):
        fixture = fixtures.PipelineIntegrationTests()
        fixture.setUp()
        try:
            fixture._write_csv()
            self.assertEqual(fixture._run().processed, 1)
            before = deepcopy(read_snapshot(fixture.product)["wafers"])
            with patch("atlas_pipeline.pipeline.merge_and_clean", side_effect=AssertionError("no recalculation")):
                result = fixture._run(rebuild_reports=True)
            self.assertEqual((result.processed, len(result.reports)), (0, 3))
            self.assertEqual(before, read_snapshot(fixture.product)["wafers"])
            for filename in result.reports:
                with Path(filename).open("rb") as stream:
                    book = load_workbook(stream)
                    try:
                        for sheet in book:
                            for row in sheet:
                                for cell in row:
                                    self.assertEqual((cell.alignment.horizontal, cell.alignment.vertical),
                                                     ("center", "center"), (sheet.title, cell.coordinate))
                        self.assertLess(book.worksheets[0].column_dimensions["C"].width, 14)
                    finally:
                        book.close()
            current = list(before.values())
            for builder in (_bin_book, _test_book):
                cached = deepcopy(current)
                book = builder(current, current, [])
                book.close()
                self.assertEqual(current, cached)
        finally:
            fixture.tearDown()


if __name__ == "__main__":
    unittest.main()
