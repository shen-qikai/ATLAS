"""Three Excel projections of cached Wafer metrics, never of historical raw CSVs."""

from functools import lru_cache
import math
import os
from pathlib import Path
import shutil
import tempfile
import unicodedata

from PIL import ImageFont

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COMMON_COLUMNS = [
    "Product", "Lot", "Wafer", "有效die数量", "晶圆总die数量", "记录数", "BIN有效记录数",
    "期望die数量", "die差值", "数量状态", "数量统计方式", "规格版本", "更新时间",
]


def _common(record):
    metrics = record["metrics"]
    return [
        record["product"], record["lot"], "W" + record["wafer"].zfill(2),
        metrics["effective_die_count"], metrics.get("wafer_total_die_count", metrics["record_count"]),
        metrics["record_count"], metrics["valid_bin_count"],
        metrics["expected_die_count"], metrics["die_delta"], metrics["die_count_status"],
        metrics["count_method"], metrics["spec_version"], record["updated_at"],
    ]


@lru_cache(maxsize=8)
def _sizing_font(points, bold=False, normal=False):
    """Use actual Windows fonts when available; no Excel process is required."""
    fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    filename = "calibri.ttf" if normal else ("msyhbd.ttc" if bold else "msyh.ttc")
    try:
        return ImageFont.truetype(str(fonts / filename), round(points * 96 / 72),
                                  layout_engine=ImageFont.Layout.BASIC)
    except (OSError, AttributeError):
        return None


def _display_text(cell):
    """Measure displayed precision, not Python's full float representation."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if cell.number_format == "0.00%":
            return f"{value:.2%}"
        if cell.number_format == "0.000000":
            return f"{value:.6f}"
        if cell.number_format == "0":
            return f"{value:.0f}"
    return str(value)


@lru_cache(maxsize=8192)
def _text_pixels(text, points=10, bold=False):
    font = _sizing_font(points, bold)
    if font is not None:
        return max((font.getlength(line) for line in text.split("\n")), default=0)
    # Chinese/full-width characters need more room than Latin characters.
    return max((sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)
                * points * 96 / 72 * 0.55 for line in text.split("\n")), default=0)


def _fit_columns(sheet):
    normal = _sizing_font(11, normal=True)
    digit_width = normal.getlength("0") if normal is not None else 7
    for index, column in enumerate(sheet.iter_cols(), start=1):
        widths = [_text_pixels(_display_text(cell), cell.font.sz or 11, bool(cell.font.bold))
                  for cell in column]
        pixels = max(widths, default=0)
        width = max(1, math.ceil((pixels + 6) / digit_width * 256) / 256)
        sheet.column_dimensions[get_column_letter(index)].width = min(255, width)
        # Excel has a maximum column width. Wrap only those exceptional long cells.
        for cell, cell_width in zip(column, widths):
            text = _display_text(cell)
            available = 255 * digit_width - 6
            if (width > 255 and cell_width > available) or "\n" in text:
                lines = sum(max(1, math.ceil(_text_pixels(
                    line, cell.font.sz or 11, bool(cell.font.bold)) / available))
                    for line in text.split("\n"))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                height = max(15, (cell.font.sz or 11) * 1.5) * max(1, lines)
                sheet.row_dimensions[cell.row].height = min(409, max(
                    sheet.row_dimensions[cell.row].height or sheet.sheet_format.defaultRowHeight or 15,
                    height))


def _style(sheet, data_start=2, percent_columns=(), numeric_columns=()):
    sheet.freeze_panes = f"D{data_start}"
    sheet.sheet_format.defaultRowHeight = 15
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="235789")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.row < data_start:
                cell.fill = PatternFill("solid", fgColor="E5EDF5")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F6FA")
            if cell.row >= data_start and isinstance(cell.value, (int, float)):
                if cell.column in percent_columns:
                    cell.number_format = "0.00%"
                elif cell.column in numeric_columns:
                    cell.number_format = "0.000000"
                else:
                    cell.number_format = "0"
    sheet.row_dimensions[1].height = 20
    _fit_columns(sheet)
    if data_start == 2:
        sheet.auto_filter.ref = sheet.dimensions


def _metadata_rows(sheet, records, columns):
    for field, label in (("unit", "单位"), ("lower", "下限"), ("upper", "上限")):
        row = [label] + [None] * (len(COMMON_COLUMNS) - 1)
        for column in columns:
            values = {record["metrics"]["tests"][column][field]
                      for record in records if column in record["metrics"]["tests"]}
            row.append(next(iter(values)) if len(values) == 1 else "见测试项规格")
        sheet.append(row)


def _extra_sheets(book, records, scan_errors, with_specs=False):
    status = book.create_sheet("处理状态")
    status.append(["Product", "Lot", "Wafer", "状态", "错误或警告", "更新时间", "原始文件"])
    for record in records:
        warning = record.get("error", "") or "; ".join(record.get("metrics", {}).get("warnings", []))
        status.append([
            record["product"], record["lot"], "W" + record["wafer"].zfill(2),
            record["status"], warning, record.get("attempted_at", record.get("updated_at", "")),
            "; ".join(source["relative_path"] for source in record.get("source_files", [])),
        ])
    for error in scan_errors:
        status.append([None, None, None, "VALIDATION_FAILED", error])
    _style(status)

    if with_specs:
        specs = book.create_sheet("测试项规格")
        specs.append(["规格版本", "测试项", "单位", "下限", "上限"])
        emitted = set()
        for record in records:
            if record.get("status") != "current":
                continue
            metrics = record["metrics"]
            for column, test in metrics["tests"].items():
                key = (metrics["spec_version"], column)
                if key not in emitted:
                    specs.append([key[0], column, test["unit"], test["lower"], test["upper"]])
                    emitted.add(key)
        _style(specs, numeric_columns=(4, 5))

    info = book.create_sheet("数据说明")
    info.append(["项目", "说明"])
    info.append(["来源", "来自 .atlas/state.sqlite 指标缓存；不会重新读取全部历史原始数据"])
    info.append(["纳入范围", "仅 current 状态；failed/missing 的旧指标保留在缓存但不纳入主表"])
    info.append(["总良率", "默认 Good BIN = 1；分母为清理后的记录数，与原 BIN 占比口径一致"])
    info.append(["有效die数量", "优先统计有效唯一坐标；无坐标时为有效BIN记录数，数量统计方式列注明"])
    info.append(["数量状态", "期望值/容差由产品配置指定；偏差仅警告，不证明漏测或完整性"])
    info.append(["测试项失效率", "超规格失效die数/本片总die数；坐标可识别时按每颗die最后记录统计，不按有效测试值数或期望die数作分母"])
    info.append(["缺测", "空白、文本及非有限测试值单列缺测数量，不计入超规格失效；等于限值算合格"])
    info.append(["数据标识", "不从文件名猜测CP/FT；一个产品/Lot/Wafer只存放一套数据，RT仍按轮次整合"])
    info.append(["测试项平均值", "所有有效有限数值的均值；文本列或无有效数据为 NA"])
    info.append(["规格变更", "不同规格保留独立版本；主表元信息冲突时请查测试项规格 Sheet"])
    _style(info)


def _bin_book(current, all_records, scan_errors):
    book = Workbook()
    sheet = book.active
    sheet.title = "BIN汇总"
    bins = {key for record in current for key in record["metrics"]["bin_counts"]}
    bins = sorted(bins, key=lambda value: (0, int(value)) if value.isdigit() else (1, value))
    sheet.append(COMMON_COLUMNS + ["总良率"] + [f"BIN_{value}" for value in bins])
    for record in current:
        metrics = record["metrics"]
        total = metrics["record_count"]
        sheet.append(_common(record) + [metrics["good_count"] / total] + [
            metrics["bin_counts"].get(value, 0) / total for value in bins
        ])
    _style(sheet, percent_columns=range(len(COMMON_COLUMNS) + 1, sheet.max_column + 1))
    _extra_sheets(book, all_records, scan_errors)
    return book


def _test_book(current, all_records, scan_errors, mean=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "测试项均值" if mean else "测试项失效率"
    columns = list(dict.fromkeys(column for record in current for column in record["metrics"]["tests"]))
    sheet.append(COMMON_COLUMNS + columns)
    _metadata_rows(sheet, current, columns)
    for record in current:
        metrics = record["metrics"]
        row = _common(record)
        for column in columns:
            test = metrics["tests"].get(column)
            if not test:
                value = "NA"
            elif mean:
                value = test["mean"] if test["mean"] is not None else "NA"
            elif test["passed_count"] is None:
                value = "No_SPEC"
            else:
                total, failed, missing = _die_counts(metrics, test)
                value = failed / total if total else "No_data"
            row.append(value)
        sheet.append(row)
    numeric = range(len(COMMON_COLUMNS) + 1, sheet.max_column + 1)
    _style(sheet, data_start=5, percent_columns=() if mean else numeric,
           numeric_columns=numeric if mean else ())
    counts = book.create_sheet("测试项有效数")
    counts.append(COMMON_COLUMNS + columns)
    for record in current:
        counts.append(_common(record) + [
            record["metrics"]["tests"].get(column, {}).get("valid_count", "NA") for column in columns
        ])
    _style(counts)
    if not mean:
        for title, field in (("测试项失效die数", "failed"), ("测试项缺测die数", "missing")):
            detail = book.create_sheet(title)
            detail.append(COMMON_COLUMNS + columns)
            for record in current:
                metrics = record["metrics"]
                total = metrics.get("wafer_total_die_count", metrics["record_count"])
                values = []
                for column in columns:
                    test = metrics["tests"].get(column)
                    if test is None:
                        values.append("NA")
                    elif field == "failed" and test["passed_count"] is None:
                        values.append("No_SPEC")
                    else:
                        _, failed, missing = _die_counts(metrics, test)
                        values.append(failed if field == "failed" else missing)
                detail.append(_common(record) + values)
            _style(detail)
    _extra_sheets(book, all_records, scan_errors, with_specs=True)
    return book


def _die_counts(metrics, test):
    total = metrics.get("wafer_total_die_count", metrics["record_count"])
    if "wafer_total_die_count" not in metrics and metrics["record_count"] != metrics["effective_die_count"]:
        raise ValueError("旧缓存含重复或不可识别die，请先运行增量处理升级指标，再重建失效率报表")
    failed = test.get("failed_die_count", None)
    if failed is None and test["passed_count"] is not None:
        failed = test["valid_count"] - test["passed_count"]
    missing = test.get("missing_die_count", total - test["valid_count"])
    return total, failed, missing


def export_reports(product_folder: Path, wafer_records: dict, scan_errors=None) -> list[Path]:
    records = sorted(wafer_records.values(), key=lambda record: (
        record["lot"].casefold(), int(record["wafer"]), record["stage"]
    ))
    current = [record for record in records if record.get("status") == "current"]
    errors = scan_errors or []
    paths = [product_folder / f"{product_folder.name}_{name}.xlsx"
             for name in ("BIN良率", "测试项失效率", "测试项平均值")]
    if any(path.is_symlink() or product_folder not in path.resolve().parents for path in paths):
        raise ValueError("报表路径不得指向产品目录之外")
    builders = [lambda: _bin_book(current, records, errors),
                lambda: _test_book(current, records, errors),
                lambda: _test_book(current, records, errors, mean=True)]
    with tempfile.TemporaryDirectory(dir=product_folder / ".atlas", prefix="reports_") as temporary:
        staging = Path(temporary)
        backups, installed = {}, []
        for index, builder in enumerate(builders):
            book = builder()
            try:
                book.save(staging / f"{index}.xlsx")
            finally:
                book.close()
        for index, path in enumerate(paths):
            if path.exists():
                backup = staging / f"{index}.previous.xlsx"
                shutil.copy2(path, backup)
                backups[path] = backup
        try:
            for index, path in enumerate(paths):
                os.replace(staging / f"{index}.xlsx", path)
                installed.append(path)
        except OSError as exc:
            rollback_errors = []
            for path in reversed(installed):
                try:
                    if path in backups:
                        os.replace(backups[path], path)
                    else:
                        path.unlink()
                except OSError as rollback_error:
                    rollback_errors.append(str(rollback_error))
            detail = f"报表更新失败，请关闭 Excel 后重试。指标已缓存，不需重算 Wafer: {exc}"
            if rollback_errors:
                detail += "; 报表回滚异常: " + "; ".join(rollback_errors)
            raise RuntimeError(detail) from exc
    return paths
